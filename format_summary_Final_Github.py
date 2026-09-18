#!/usr/bin/env python3
"""Render summary tables as a colour-coded workbook.
"""

import os
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- run config
# Set SOURCE and run this file. Everything else is optional.
#
#   SOURCE = "summary_flip_all_g_p5_s4631010000_LogReg.xlsx"   -> every sheet
#   SOURCE = "summary_flip_g2_p50_s392337000_LogReg.csv"       -> one sheet
#
# Command-line arguments, if given, override all three of these.

SOURCE = "summary_flip_all_g_p5_s432187000.xlsx"      # input file: .xlsx (all sheets) or .csv (one sheet)
OUTPUT = ""      # output path; blank -> "<source stem>_formatted.xlsx"
LABEL = ""       # optional free text for the first legend line
# -----------------------------------------------------------------------------

REFERENCE = "logistic"
LOWER_IS_BETTER = {"MSE_true", "Brier", "LogLoss"}   # Accuracy, AUC: higher

BLUE = PatternFill("solid", fgColor="4472C4")     # block changed
RED = PatternFill("solid", fgColor="C00000")      # new metric
L_RED = PatternFill("solid", fgColor="FFC7CE")    # the logistic reference
L_YELLOW = PatternFill("solid", fgColor="FFEB9C") # significantly better
L_GREEN = PatternFill("solid", fgColor="C6EFCE")  # degrades less
HEADER = PatternFill("solid", fgColor="D9D9D9")

FONT = "Arial"
NUMERIC = {"mean", "se", "diff_vs_logistic", "diff_se",
           "clean_mean", "increase", "increase_se", "tau", "sd"}

# DGP internals: context for the run, never a column anyone reads across.
# Dropped from the table and recorded in the legend instead. bayes and rho are
# kept even when constant, since they are design parameters meant to be swept.
DROP = ["base_seed", "g", "p", "prevalence", "tau", "sd", "dataset"]


def better(metric, diff):
    if pd.isna(diff):
        return False
    return diff < 0 if metric in LOWER_IS_BETTER else diff > 0


def more_robust(metric, inc, ref):
    """Less degradation than logistic, in the direction that helps."""
    if pd.isna(inc) or ref is None or pd.isna(ref):
        return False
    return inc < ref if metric in LOWER_IS_BETTER else inc > ref


def fmt(v):
    """Render a dropped constant for the legend line."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "n/a"
    return f"{v:.6g}" if isinstance(v, float) else str(v)


def strip_constants(df):
    """Drop the DGP context columns, after asserting each really is constant."""
    dropped = {}
    for c in DROP:
        if c not in df.columns:
            continue
        assert df[c].nunique(dropna=False) == 1, f"{c} varies; refusing to drop it"
        dropped[c] = df[c].iloc[0]
    return df.drop(columns=list(dropped)), dropped


def render(ws, df):
    """Write one summary frame into one worksheet."""
    cols = list(df.columns)
    ncol = len(cols)
    idx = {c: i + 1 for i, c in enumerate(cols)}

    block_cols = [c for c in ("contamination", "flip_rule", "share") if c in cols]
    key = df[block_cols].astype(object).where(df[block_cols].notna(), "")

    ref = df[df["model"] == REFERENCE]
    ref_inc = {(m, *k): v for m, k, v in zip(
        ref["metric"],
        key.loc[ref.index].itertuples(index=False, name=None),
        ref["increase"])}

    thin = Side(style="thin", color="BFBFBF")
    for j, name in enumerate(cols, start=1):
        c = ws.cell(1, j, name)
        c.font = Font(name=FONT, size=10, bold=True)
        c.fill = HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="medium", color="808080"))

    def band(row, fill, label):
        for j in range(1, ncol + 1):
            c = ws.cell(row, j, label if j == 1 else None)
            c.fill = fill
            c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        ws.row_dimensions[row].height = 15

    r = 2
    prev_metric = prev_block = None
    counts = dict(red=0, blue=0, lred=0, yellow=0, green=0)

    for i, rec in df.iterrows():
        metric = rec["metric"]
        block = tuple(key.loc[i])

        if metric != prev_metric:
            # a new metric also opens a new block; draw one band, not two
            band(r, RED, str(metric))
            counts["red"] += 1
            r += 1
        elif block != prev_block:
            parts = [f"{v:.0%}" if isinstance(v, float) else str(v)
                     for v in block if v != ""]
            band(r, BLUE, "  ".join(parts))
            counts["blue"] += 1
            r += 1
        prev_metric, prev_block = metric, block

        for j, col in enumerate(cols, start=1):
            v = rec[col]
            c = ws.cell(r, j, None if pd.isna(v) else v)
            c.font = Font(name=FONT, size=10)
            c.border = Border(bottom=thin)
            if col in NUMERIC:
                c.number_format = "0.000000"
                c.alignment = Alignment(horizontal="right")
            elif col in ("share", "bayes", "rho"):
                c.number_format = "0%" if col == "share" else "0.00"
                c.alignment = Alignment(horizontal="center")
            elif col == "significant":
                c.number_format = "0"
                c.alignment = Alignment(horizontal="center")

        if rec["model"] == REFERENCE:
            for col in ("model", "mean", "increase"):
                if col == "increase" and pd.isna(rec["increase"]):
                    continue
                ws.cell(r, idx[col]).fill = L_RED
                counts["lred"] += 1
        else:
            if rec["significant"] == 1 and better(metric, rec["diff_vs_logistic"]):
                for col in ("model", "mean", "diff_vs_logistic"):
                    ws.cell(r, idx[col]).fill = L_YELLOW
                counts["yellow"] += 1
            if more_robust(metric, rec["increase"], ref_inc.get((metric, *block))):
                ws.cell(r, idx["increase"]).fill = L_GREEN
                counts["green"] += 1
        r += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}1"
    widths = {"metric": 11, "contamination": 14, "flip_rule": 11, "share": 8,
              "model": 16, "bayes": 8, "rho": 7}
    for j, col in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(col, 15)

    return counts, block_cols, r - 1


def load(spec):
    """Read a source given as path.csv or path.xlsx:sheet."""
    if ".xlsx:" in spec:
        path, sheet = spec.rsplit(":", 1)
        return pd.read_excel(path, sheet_name=sheet), f"{path.split('/')[-1]}[{sheet}]"
    return pd.read_csv(spec), spec.split("/")[-1]


def sheet_name_for(df, path):
    """Name a worksheet from the data itself, falling back to the file stem."""
    for col in ("g", "dataset"):
        if col in df.columns and df[col].nunique(dropna=False) == 1:
            v = df[col].iloc[0]
            return f"g{v}" if col == "g" else str(v)[:31]
    return os.path.splitext(os.path.basename(path))[0][:31]


def expand(source):
    """Turn one input path into the (sheet_name, spec) pairs build() wants."""
    if source.lower().endswith((".xlsx", ".xlsm")):
        names = pd.ExcelFile(source).sheet_names
        return [(n, f"{source}:{n}") for n in names]
    df, _ = load(source)
    return [(sheet_name_for(df, source), source)]


def default_output(source):
    return os.path.splitext(source)[0] + "_formatted.xlsx"


def build(dst, sources, label=""):
    wb = Workbook()
    wb.remove(wb.active)

    notes = []
    for name, path in sources:
        df, srcname = load(path)
        df, dropped = strip_constants(df)
        ws = wb.create_sheet(name)
        counts, block_cols, total = render(ws, df)
        notes.append((name, srcname, len(df), counts, dropped, block_cols))
        print(f"{name}: data rows {len(df)} | red {counts['red']} | "
              f"blue {counts['blue']} | total rows {total}")
        print(f"  highlights: light red {counts['lred']}, "
              f"yellow {counts['yellow']}, green {counts['green']}")
        print(f"  dropped: {dropped}")

    lg = wb.create_sheet("legend")
    lg.column_dimensions["A"].width = 4
    lg.column_dimensions["B"].width = 110
    bc = notes[0][5]
    entries = [
        (None, "Legend", True),
        (None, (f"{label} " if label else "") + "One sheet per signal function; "
               f"rows ordered metric > {' > '.join(bc)} > loss function.", False),
        (None, "", False),
    ]
    for name, path, n, counts, dropped, _ in notes:
        ctx = ", ".join(f"{k} = {fmt(v)}" for k, v in dropped.items())
        entries.append(
            (None, f"Sheet '{name}': {n} rows, source {path}. Constant context: "
                   f"{ctx}. Highlights: {counts['yellow']} yellow, "
                   f"{counts['green']} green.", False))
    entries += [
        (None, "", False),
        (None, "Columns dropped from the tables because they are constant within a "
               "sheet and describe the run rather than a swept factor: "
               f"{', '.join(notes[0][4])}. Their values are recorded above. "
               "bayes and rho, where present, are kept even when constant, since "
               "they are design parameters meant to be swept.", False),
        (None, "", False),
        (RED, "New metric begins.", False),
        (BLUE, "New " + " / ".join(bc) + " block. A new metric also opens a "
               "block, so only the red band is drawn there.", False),
        (None, "", False),
        (L_RED, f"The {REFERENCE} reference: model name, mean, and increase over "
                "clean where applicable.", False),
        (L_YELLOW, "Significantly better than logistic in this block: name, mean, "
                   "difference. The file's own two-sided flag (|diff| > 2 SE of the "
                   "paired difference) combined with the sign that means improvement "
                   "- lower for MSE_true / Brier / LogLoss, higher for Accuracy / AUC.",
         False),
        (L_GREEN, "Degrades less than logistic under the same contamination: a "
                  "smaller increase for MSE_true / Brier / LogLoss, a less negative "
                  "one for Accuracy / AUC. A direct comparison of two means, not a "
                  "test - the file has no SE for the difference of paired differences.",
         False),
    ]
    for i, (fill, text, bold) in enumerate(entries, start=1):
        if fill is not None:
            lg.cell(i, 1).fill = fill
        c = lg.cell(i, 2, text)
        c.font = Font(name=FONT, size=10, bold=bold)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if len(text) > 108:
            lg.row_dimensions[i].height = 15 * (len(text) // 108 + 1)

    wb.save(dst)
    print(f"wrote {dst}")


def main():
    args = sys.argv[1:]

    if not args:
        if not SOURCE:
            print(__doc__)
            print("Nothing to do: set SOURCE at the top of this file, "
                  "or pass arguments on the command line.")
            return
        build(OUTPUT or default_output(SOURCE), expand(SOURCE), LABEL)
        return

    if args[0] in ("-h", "--help"):
        print(__doc__)
        return
    label = ""
    if args and args[0].startswith("--label="):
        label = args.pop(0).split("=", 1)[1]
    dst = args[0]
    srcs = [tuple(a.split("=", 1)) for a in args[1:]]
    build(dst, srcs, label)


if __name__ == "__main__":
    main()
